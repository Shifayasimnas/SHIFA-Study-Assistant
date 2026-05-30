import html
import hashlib
import io
import re
import sqlite3
import zipfile
from concurrent.futures import ThreadPoolExecutor, as_completed
from contextlib import contextmanager
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET

import PyPDF2
import streamlit as st
from openai import OpenAI
from PIL import Image

TESSERACT_EXE = Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe")

# =========================================================
# APP CONFIG
# =================================s========================

APP_TITLE = "SHIFA Study Assistant"
DB_PATH = Path("history.db")

# WORKING OPENROUTER MODEL
DEFAULT_MODEL = "openai/gpt-4o-mini"
FALLBACK_MODEL = "openai/gpt-4o-mini"

st.set_page_config(
    page_title=APP_TITLE,
    page_icon="🤖",
    layout="wide",
    initial_sidebar_state="expanded",
)

# =========================================================
# THEME
# =========================================================

def inject_theme():
    st.markdown("""
    <style>

    .stApp {
        background:
        radial-gradient(circle at top left, rgba(34,211,238,0.15), transparent 30%),
        linear-gradient(135deg,#020617 0%,#07111f 50%,#081827 100%);
        color: #cbd5e1;
    }

    .stApp,
    .stMarkdown,
    [data-testid="stMarkdownContainer"] {
        color:#cbd5e1;
    }

    section[data-testid="stSidebar"] {
        background: #020617;
        border-right: 1px solid rgba(125,249,255,0.15);
    }

    section[data-testid="stSidebar"] * {
        color: white !important;
    }

    h1, h2, h3{
        color:#f8fafc !important;
        letter-spacing:0;
        text-shadow:0 2px 16px rgba(100,255,218,0.16);
    }

    p{
        color:#cbd5e1;
        line-height:1.75;
    }

    label,
    [data-testid="stWidgetLabel"],
    [data-testid="stWidgetLabel"] p{
        color:#f8fafc !important;
        font-weight:700 !important;
        text-shadow:0 1px 10px rgba(100,255,218,0.12);
    }

    .stCaption,
    [data-testid="stCaptionContainer"],
    [data-testid="stCaptionContainer"] p{
        color:#a7f3d0 !important;
    }

    .main-title,
    .page-title{
        font-size:60px;
        font-weight:900;
        color:#f8fafc;
        margin-bottom:10px;
        line-height:1.05;
        letter-spacing:0;
        text-shadow:
            0 2px 18px rgba(100,255,218,0.24),
            0 0 34px rgba(34,211,238,0.12);
    }

    .section-title{
        color:#64ffda;
        font-size:28px;
        font-weight:800;
        margin:22px 0 12px;
        text-shadow:0 2px 16px rgba(100,255,218,0.18);
    }

    .sub-title{
        font-size:18px;
        color:#cbd5e1;
        margin-bottom:30px;
        line-height:1.7;
    }

    .glass-card,
    .feature-card{
        background: rgba(15,23,42,0.65);
        border:1px solid rgba(125,249,255,0.22);
        border-radius:18px;
        padding:22px;
        backdrop-filter: blur(14px);
        transition:transform 0.25s ease, box-shadow 0.25s ease, border-color 0.25s ease;
        min-height:180px;
        color:#cbd5e1;
    }

    .glass-card:hover,
    .feature-card:hover{
        transform:translateY(-4px);
        border-color:rgba(100,255,218,0.46);
        box-shadow:
            0 0 26px rgba(34,211,238,0.18),
            0 18px 40px rgba(2,6,23,0.32);
    }

    .glass-card h3,
    .feature-card h3{
        color:#64ffda !important;
        margin-bottom:10px;
        text-shadow:0 2px 14px rgba(100,255,218,0.2);
    }

    .glass-card p,
    .feature-card p{
        color:#cbd5e1;
        line-height:1.7;
    }

    .chat-user,
    .chat-bubble-user{
        background:linear-gradient(135deg,#2563eb,#0891b2);
        padding:18px;
        border-radius:16px;
        margin:12px 0;
        color:#f8fafc;
        border:1px solid rgba(248,250,252,0.14);
        box-shadow:0 14px 30px rgba(8,145,178,0.16);
    }

    .chat-user b,
    .chat-bubble-user b{
        color:#ffffff;
        text-shadow:0 1px 12px rgba(255,255,255,0.18);
    }

    .chat-ai,
    .chat-bubble-model{
        background:rgba(15,23,42,0.8);
        border:1px solid rgba(125,249,255,0.25);
        padding:18px;
        border-radius:16px;
        margin:12px 0;
        color:#cbd5e1;
        box-shadow:0 14px 32px rgba(2,6,23,0.28);
    }

    .chat-ai b,
    .chat-bubble-model b{
        color:#64ffda;
        text-shadow:0 1px 12px rgba(100,255,218,0.18);
    }

    .stTextInput input,
    .stTextArea textarea,
    .stSelectbox div[data-baseweb="select"] > div,
    .stSlider [data-baseweb="slider"]{
        background:#081827 !important;
        color:#f8fafc !important;
        border:1px solid rgba(125,249,255,0.34) !important;
        border-radius:14px !important;
        box-shadow:0 0 0 1px rgba(2,6,23,0.18);
    }

    .stTextInput input:focus,
    .stTextArea textarea:focus{
        border-color:rgba(100,255,218,0.68) !important;
        box-shadow:0 0 0 1px rgba(100,255,218,0.22), 0 0 20px rgba(34,211,238,0.12) !important;
    }

    .stTextInput input::placeholder,
    .stTextArea textarea::placeholder{
        color:#94a3b8 !important;
        opacity:1 !important;
    }

    .stSelectbox,
    .stSlider,
    .stTextInput,
    .stTextArea{
        color:#f8fafc;
    }

    .stSelectbox *{
        color:#f8fafc !important;
    }

    .stSlider [data-testid="stTickBar"] *{
        color:#cbd5e1 !important;
    }

    .stButton > button{
        background:linear-gradient(135deg,#67e8f9,#22d3ee);
        color:black;
        border:none;
        border-radius:14px;
        font-weight:700;
        height:3rem;
    }

    .stButton > button:hover{
        box-shadow:0 0 25px rgba(34,211,238,0.3);
    }

    .history-card,
    .empty-history{
        background:rgba(15,23,42,0.7);
        border:1px solid rgba(125,249,255,0.24);
        padding:20px;
        border-radius:16px;
        margin-bottom:15px;
        color:#cbd5e1;
        transition:transform 0.25s ease, box-shadow 0.25s ease, border-color 0.25s ease;
    }

    .history-card:hover,
    .empty-history:hover{
        transform:translateY(-2px);
        border-color:rgba(100,255,218,0.44);
        box-shadow:0 0 24px rgba(34,211,238,0.14);
    }

    .history-card h4,
    .empty-history h4{
        color:#64ffda;
        font-weight:800;
        text-shadow:0 2px 14px rgba(100,255,218,0.16);
    }

    .history-card p,
    .empty-history p{
        color:#cbd5e1;
        line-height:1.75;
    }

    .history-time{
        color:#a7f3d0;
        font-size:13px;
    }

    .upload-hero{
        background:
            linear-gradient(135deg, rgba(15,23,42,0.86), rgba(8,47,73,0.58)),
            radial-gradient(circle at 84% 10%, rgba(34,211,238,0.22), transparent 28%);
        border:1px solid rgba(125,249,255,0.26);
        border-radius:18px;
        padding:26px;
        margin:8px 0 22px;
        box-shadow:0 24px 70px rgba(2,6,23,0.36);
    }

    .upload-hero h2{
        margin:0 0 8px;
        color:#f8fafc !important;
    }

    .upload-hero p{
        margin:0;
        color:#cbd5e1;
    }

    .file-card{
        background:rgba(15,23,42,0.72);
        border:1px solid rgba(125,249,255,0.22);
        border-radius:14px;
        padding:16px;
        margin-bottom:12px;
        box-shadow:0 12px 32px rgba(2,6,23,0.24);
    }

    .file-row{
        display:flex;
        align-items:flex-start;
        gap:14px;
    }

    .file-icon{
        width:46px;
        height:46px;
        border-radius:12px;
        display:flex;
        align-items:center;
        justify-content:center;
        color:#020617;
        font-size:22px;
        font-weight:900;
        background:linear-gradient(135deg,#67e8f9,#a7f3d0);
        flex:0 0 auto;
    }

    .file-name{
        color:#f8fafc;
        font-size:16px;
        font-weight:800;
        word-break:break-word;
    }

    .file-meta{
        color:#94a3b8;
        font-size:13px;
        margin-top:5px;
    }

    .status-pill{
        display:inline-flex;
        padding:4px 10px;
        border-radius:999px;
        margin-top:10px;
        font-size:12px;
        font-weight:800;
        color:#082f49;
        background:#67e8f9;
    }

    .metric-strip{
        display:grid;
        grid-template-columns:repeat(4, minmax(0, 1fr));
        gap:12px;
        margin:14px 0 20px;
    }

    .study-metric{
        background:rgba(8,24,39,0.76);
        border:1px solid rgba(125,249,255,0.2);
        border-radius:14px;
        padding:14px;
    }

    .study-metric strong{
        color:#67e8f9;
        font-size:24px;
        display:block;
    }

    .study-metric span{
        color:#cbd5e1;
        font-size:13px;
    }

    .processing-ring{
        width:28px;
        height:28px;
        border:3px solid rgba(103,232,249,0.22);
        border-top-color:#67e8f9;
        border-radius:50%;
        animation:spin 0.9s linear infinite;
        display:inline-block;
        vertical-align:middle;
        margin-right:10px;
    }

    @keyframes spin{
        to{ transform:rotate(360deg); }
    }

    @media (max-width: 760px){
        .metric-strip{
            grid-template-columns:repeat(2, minmax(0, 1fr));
        }
        .main-title,
        .page-title{
            font-size:42px;
        }
    }

    </style>
    """, unsafe_allow_html=True)

# =========================================================
# DATABASE
# =========================================================

@contextmanager
def db_connection():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    try:
        yield conn
        conn.commit()
    finally:
        conn.close()

def init_db():
    with db_connection() as conn:
        conn.execute("""
        CREATE TABLE IF NOT EXISTS history(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            question TEXT,
            answer TEXT,
            time TEXT
        )
        """)

def save_history(question, answer):
    with db_connection() as conn:
        conn.execute(
            "INSERT INTO history(question,answer,time) VALUES(?,?,?)",
            (
                question,
                answer,
                datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            )
        )

def get_history(search=""):
    with db_connection() as conn:

        if search:
            return conn.execute("""
            SELECT * FROM history
            WHERE question LIKE ?
            ORDER BY id DESC
            """, (f"%{search}%",)).fetchall()

        return conn.execute("""
        SELECT * FROM history
        ORDER BY id DESC
        """).fetchall()

def clear_history():
    with db_connection() as conn:
        conn.execute("DELETE FROM history")

# =========================================================
# OPENROUTER CLIENT
# =========================================================

def get_client():

    api_key = st.secrets.get("OPENROUTER_API_KEY", "")

    if not api_key:
        return None

    return OpenAI(
        api_key=api_key,
        base_url="https://openrouter.ai/api/v1",
        timeout=60,
        default_headers={
            "HTTP-Referer": "http://localhost:8501",
            "X-Title": APP_TITLE,
        }
    )

# =========================================================
# AI CHAT
# =========================================================

def ai_chat(messages, temperature=0.3):

    client = get_client()

    if client is None:
        return "❌ OpenRouter API key missing."

    try:

        response = client.chat.completions.create(
            model="openai/gpt-4o-mini",
            messages=messages,
            temperature=temperature,
            max_tokens=1800
        )

        return response.choices[0].message.content.strip()

    except Exception as e:

        return f"""
❌ AI Connection Error

Possible reasons:
- Internet issue
- Wrong API key
- OpenRouter server issue
- Firewall/VPN blocking connection

Error:
{str(e)}
"""

# =========================================================
# HELPERS
# =========================================================

def esc(text):
    return html.escape(str(text))

def render_title(title):
    st.markdown(f"""
    <div class="main-title page-title">{title}</div>
    """, unsafe_allow_html=True)

# =========================================================
# SIDEBAR
# =========================================================

def sidebar():

    with st.sidebar:

        st.image(
            "https://cdn-icons-png.flaticon.com/512/4712/4712109.png",
            width=120
        )

        st.markdown("## SHIFA Study Assistant")
        st.caption("Your Personal AI Engineering Companion")

        menu = st.radio(
            "Navigation",
            [
                "Dashboard",
                "AI Study Chat",
                "Upload Center",
                "MCQ Generator",
                "Search History"
            ]
        )

        st.divider()

        if st.button("🆕 New Chat", use_container_width=True):
            st.session_state.messages = []
            st.rerun()

    return menu

# =========================================================
# DASHBOARD
# =========================================================

def dashboard():

    render_title("SHIFA Study Assistant")

    st.markdown("""
    <div class="sub-title">
    Premium AI-powered engineering learning workspace
    </div>
    """, unsafe_allow_html=True)

    col1, col2, col3 = st.columns(3)

    with col1:
        st.markdown("""
        <div class="glass-card feature-card">
        <h3>📘 AI Study Chat</h3>
        <p>
        Ask SHIFA AI anything related to engineering,
        coding, exams, AI/ML, programming, or semester preparation.
        </p>
        </div>
        """, unsafe_allow_html=True)

    with col2:
        st.markdown("""
        <div class="glass-card feature-card">
        <h3>📄 PDF Upload</h3>
        <p>
        Upload lecture notes, textbooks, PDFs, and study materials
        for instant AI summaries and explanations.
        </p>
        </div>
        """, unsafe_allow_html=True)

    with col3:
        st.markdown("""
        <div class="glass-card feature-card">
        <h3>🧠 MCQ Generator</h3>
        <p>
        Generate smart MCQs with answer keys,
        difficulty selection, and exam-focused preparation.
        </p>
        </div>
        """, unsafe_allow_html=True)

# =========================================================
# CHAT PAGE
# =========================================================

def chat_page():

    render_title("AI Study Chat")

    st.caption("Ask any engineering or study question")

    for msg in st.session_state.messages:

        if msg["role"] == "user":

            st.markdown(f"""
            <div class="chat-user chat-bubble-user">
            <b>YOU</b><br><br>
            {esc(msg["content"])}
            </div>
            """, unsafe_allow_html=True)

        else:

            st.markdown(f"""
            <div class="chat-ai chat-bubble-model">
            <b>SHIFA AI</b><br><br>
            {esc(msg["content"])}
            </div>
            """, unsafe_allow_html=True)

    with st.form("chat_form", clear_on_submit=True):

        question = st.text_area(
            "Ask AI",
            placeholder="Explain DBMS normalization with examples...",
            height=120
        )

        submit = st.form_submit_button(
            "Ask SHIFA AI",
            use_container_width=True
        )

    if submit and question:

        st.session_state.messages.append({
            "role": "user",
            "content": question
        })

        messages = [
            {
                "role": "system",
                "content": """
You are SHIFA,
an expert AI engineering professor.

Give:
- accurate answers
- exam-focused explanations
- examples
- formulas
- structured points
- professional responses
"""
            }
        ]

        messages.extend(st.session_state.messages)

        with st.spinner("SHIFA AI is thinking..."):

            answer = ai_chat(messages)

        st.session_state.messages.append({
            "role": "assistant",
            "content": answer
        })

        save_history(question, answer)

        st.rerun()

# =========================================================
# UPLOAD CENTER - DOCUMENT INTELLIGENCE
# =========================================================

SUPPORTED_UPLOAD_TYPES = ["pdf", "docx", "txt", "pptx", "png", "jpg", "jpeg"]
MAX_CONTEXT_CHARS = 22000

FILE_ICONS = {
    "pdf": "PDF",
    "docx": "DOC",
    "txt": "TXT",
    "pptx": "PPT",
    "png": "IMG",
    "jpg": "IMG",
    "jpeg": "IMG",
}

def file_hash(file_bytes):
    return hashlib.sha256(file_bytes).hexdigest()

def format_file_size(size_bytes):
    if size_bytes < 1024:
        return f"{size_bytes} B"
    if size_bytes < 1024 * 1024:
        return f"{size_bytes / 1024:.1f} KB"
    return f"{size_bytes / (1024 * 1024):.1f} MB"

def file_extension(filename):
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

def clean_text(text):
    text = re.sub(r"\r", "\n", text or "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()

def chunk_text(text, chunk_size=4200, overlap=450):
    text = clean_text(text)
    if len(text) <= chunk_size:
        return [text] if text else []

    chunks = []
    start = 0
    while start < len(text):
        end = min(start + chunk_size, len(text))
        chunks.append(text[start:end])
        if end == len(text):
            break
        start = max(0, end - overlap)
    return chunks

def compact_context(text, max_chars=MAX_CONTEXT_CHARS):
    text = clean_text(text)
    if len(text) <= max_chars:
        return text
    head = text[: int(max_chars * 0.68)]
    tail = text[-int(max_chars * 0.22):]
    return f"{head}\n\n[Middle content compacted for faster AI response]\n\n{tail}"

def optional_package_message(package, purpose):
    return (
        f"{package} is not installed, so {purpose} was limited. "
        f"Add it to requirements.txt for stronger extraction."
    )

@st.cache_data(show_spinner=False)
def process_pdf(file_bytes):
    try:
        reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
        if getattr(reader, "is_encrypted", False):
            try:
                reader.decrypt("")
            except Exception:
                return {"text": "", "pages": 0, "warnings": ["This PDF is encrypted and could not be opened."]}

        page_texts = []
        for index, page in enumerate(reader.pages, start=1):
            try:
                extracted = page.extract_text() or ""
                if extracted.strip():
                    page_texts.append(f"[Page {index}]\n{extracted.strip()}")
            except Exception:
                page_texts.append(f"[Page {index}]\nText extraction failed for this page.")

        text = clean_text("\n\n".join(page_texts))
        warnings = []
        if not text:
            warnings.append("No readable text was found. This may be a scanned PDF and needs OCR.")
        return {"text": text, "pages": len(reader.pages), "warnings": warnings}
    except Exception as e:
        return {"text": "", "pages": 0, "warnings": [f"PDF processing failed: {e}"]}

@st.cache_data(show_spinner=False)
def process_docx(file_bytes):
    try:
        try:
            from docx import Document

            document = Document(io.BytesIO(file_bytes))
            paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
            tables = []
            for table in document.tables:
                for row in table.rows:
                    tables.append(" | ".join(cell.text.strip() for cell in row.cells))
            text = clean_text("\n".join(paragraphs + tables))
            return {"text": text, "pages": 0, "warnings": [] if text else ["DOCX opened, but no readable text was found."]}
        except ImportError:
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as archive:
                xml = archive.read("word/document.xml")
            root = ET.fromstring(xml)
            text_nodes = [node.text for node in root.iter() if node.tag.endswith("}t") and node.text]
            text = clean_text("\n".join(text_nodes))
            return {
                "text": text,
                "pages": 0,
                "warnings": [optional_package_message("python-docx", "DOCX table and layout extraction")],
            }
    except Exception as e:
        return {"text": "", "pages": 0, "warnings": [f"DOCX processing failed: {e}"]}

@st.cache_data(show_spinner=False)
def process_txt(file_bytes):
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return {"text": clean_text(file_bytes.decode(encoding)), "pages": 0, "warnings": []}
        except UnicodeDecodeError:
            continue
    return {"text": "", "pages": 0, "warnings": ["TXT encoding could not be detected."]}

@st.cache_data(show_spinner=False)
def process_pptx(file_bytes):
    try:
        try:
            from pptx import Presentation

            presentation = Presentation(io.BytesIO(file_bytes))
            slides = []
            for slide_number, slide in enumerate(presentation.slides, start=1):
                parts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        parts.append(shape.text.strip())
                if parts:
                    slides.append(f"[Slide {slide_number}]\n" + "\n".join(parts))
            text = clean_text("\n\n".join(slides))
            return {"text": text, "pages": len(presentation.slides), "warnings": [] if text else ["PPTX opened, but no readable text was found."]}
        except ImportError:
            slides = []
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as archive:
                slide_names = sorted(name for name in archive.namelist() if name.startswith("ppt/slides/slide") and name.endswith(".xml"))
                for index, name in enumerate(slide_names, start=1):
                    root = ET.fromstring(archive.read(name))
                    text_nodes = [node.text for node in root.iter() if node.tag.endswith("}t") and node.text]
                    if text_nodes:
                        slides.append(f"[Slide {index}]\n" + "\n".join(text_nodes))
            return {
                "text": clean_text("\n\n".join(slides)),
                "pages": len(slides),
                "warnings": [optional_package_message("python-pptx", "PPTX layout extraction")],
            }
    except Exception as e:
        return {"text": "", "pages": 0, "warnings": [f"PPTX processing failed: {e}"]}

@st.cache_data(show_spinner=False)
def process_image(file_bytes, extension):
    warnings = []
    try:
        image = Image.open(io.BytesIO(file_bytes))
        width, height = image.size
        mode = image.mode
        ocr_text = ""
        try:
            import pytesseract

            if TESSERACT_EXE.exists():
                pytesseract.pytesseract.tesseract_cmd = str(TESSERACT_EXE)

            ocr_text = pytesseract.image_to_string(image)
        except ImportError:
            warnings.append(optional_package_message("pytesseract", "OCR text extraction"))
        except Exception as e:
            warnings.append(f"OCR failed: {e}")

        description = f"""
Image file ({extension.upper()})
- Dimensions: {width} x {height}
- Color mode: {mode}
- OCR text:
{ocr_text.strip() if ocr_text.strip() else "No OCR text extracted."}

Visual analysis note:
Use the chat box to ask for diagram, flowchart, table, screenshot, or handwritten-note explanation. SHIFA will answer from OCR and image metadata where available.
"""
        return {
            "text": clean_text(description),
            "pages": 1,
            "warnings": warnings,
            "image_size": (width, height),
        }
    except Exception as e:
        return {"text": "", "pages": 0, "warnings": [f"Image processing failed: {e}"]}

def process_uploaded_bytes(filename, file_bytes):
    extension = file_extension(filename)
    digest = file_hash(file_bytes)

    if extension not in SUPPORTED_UPLOAD_TYPES:
        return None, f"{filename} is not supported."

    if extension == "pdf":
        result = process_pdf(file_bytes)
    elif extension == "docx":
        result = process_docx(file_bytes)
    elif extension == "txt":
        result = process_txt(file_bytes)
    elif extension == "pptx":
        result = process_pptx(file_bytes)
    elif extension in ("png", "jpg", "jpeg"):
        result = process_image(file_bytes, extension)
    else:
        result = {"text": "", "pages": 0, "warnings": ["Unsupported file type."]}

    document = {
        "id": digest,
        "name": filename,
        "size": len(file_bytes),
        "extension": extension,
        "uploaded_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "text": result.get("text", ""),
        "pages": result.get("pages", 0),
        "warnings": result.get("warnings", []),
        "chunks": chunk_text(result.get("text", "")),
    }
    return document, None

def process_uploaded_file(uploaded_file):
    return process_uploaded_bytes(uploaded_file.name, uploaded_file.getvalue())

def combined_document_context():
    docs = st.session_state.get("uploaded_documents", {})
    sections = []
    for doc in docs.values():
        if doc.get("text"):
            sections.append(f"### Source: {doc['name']}\n{doc['text']}")
    return compact_context("\n\n".join(sections))

def render_file_list(documents):
    st.markdown("### Uploaded Files")
    if not documents:
        st.info("No files uploaded yet.")
        return

    for doc in documents.values():
        icon = FILE_ICONS.get(doc["extension"], "FILE")
        status = "Ready" if doc.get("text") else "Needs attention"
        meta = f"{doc['extension'].upper()} • {format_file_size(doc['size'])} • {doc['uploaded_at']}"
        st.markdown(f"""
        <div class="file-card">
            <div class="file-row">
                <div class="file-icon">{esc(icon)}</div>
                <div>
                    <div class="file-name">{esc(doc['name'])}</div>
                    <div class="file-meta">{esc(meta)}</div>
                    <span class="status-pill">{esc(status)}</span>
                </div>
            </div>
        </div>
        """, unsafe_allow_html=True)
        for warning in doc.get("warnings", []):
            st.warning(warning)

def render_document_metrics(documents):
    total_files = len(documents)
    total_chars = sum(len(doc.get("text", "")) for doc in documents.values())
    total_chunks = sum(len(doc.get("chunks", [])) for doc in documents.values())
    total_pages = sum(int(doc.get("pages", 0) or 0) for doc in documents.values())
    st.markdown(f"""
    <div class="metric-strip">
        <div class="study-metric"><strong>{total_files}</strong><span>Files</span></div>
        <div class="study-metric"><strong>{total_pages}</strong><span>Pages / slides</span></div>
        <div class="study-metric"><strong>{total_chunks}</strong><span>AI chunks</span></div>
        <div class="study-metric"><strong>{total_chars // 1000}K</strong><span>Text chars</span></div>
    </div>
    """, unsafe_allow_html=True)

def study_system_prompt(mode="analysis"):
    return f"""
You are SHIFA, a senior Anna University exam-oriented AI study assistant.
Answer in a Gemini/Perplexity style using clear headings, bullets, tables when useful, examples, exam notes, key takeaways, and memory tricks.
Use only the uploaded content whenever possible. If the content does not contain the answer, say what is missing and give a cautious study-oriented explanation only when it helps.
Mode: {mode}
"""

def generate_summary(context):
    if not context:
        return "No readable uploaded content is available yet."

    prompt = f"""
Analyze the uploaded study material and produce:
1. Executive Summary
2. Key Concepts
3. Important Definitions
4. Formula Extraction
5. Exam Important Points
6. Quick Revision Notes
7. Frequently Asked Questions

Uploaded content:
{context}
"""
    return ai_chat([
        {"role": "system", "content": study_system_prompt("professional document analysis")},
        {"role": "user", "content": prompt},
    ], temperature=0.2)

def generate_notes(context, note_type="Generate Notes", difficulty="Medium"):
    if not context:
        return "Upload and process at least one readable file first."

    prompt = f"""
Create: {note_type}
Difficulty: {difficulty}

Requirements:
- Anna University exam orientation
- Clear headings
- Tables where useful
- Important formulas
- 2 mark, 13 mark, and 16 mark preparation angles when relevant
- Key takeaways and memory tricks

Use this uploaded content:
{context}
"""
    return ai_chat([
        {"role": "system", "content": study_system_prompt(note_type)},
        {"role": "user", "content": prompt},
    ], temperature=0.25)

def generate_mcqs(context, difficulty="Medium"):
    if not context:
        return "Upload and process at least one readable file first."

    prompt = f"""
Generate an exam question bank from the uploaded content.

Include:
- 2 Marks Questions
- 13 Marks Questions
- 16 Marks Questions
- University Exam Questions
- Objective Questions / MCQs
- Difficulty: {difficulty}
- Answer key for MCQs
- Short answer hints for theory questions

Uploaded content:
{context}
"""
    return ai_chat([
        {"role": "system", "content": study_system_prompt("smart MCQ and question generation")},
        {"role": "user", "content": prompt},
    ], temperature=0.3)

def document_chat(question, context, history):
    if not context:
        return "I do not have readable uploaded content yet. Please upload a PDF, DOCX, TXT, PPTX, or image with OCR-readable text."

    recent_history = history[-6:]
    messages = [{"role": "system", "content": study_system_prompt("document question answering")}]
    messages.append({
        "role": "user",
        "content": f"""
Uploaded document context:
{context}

Conversation so far:
{recent_history}

Question:
{question}

Answer only from uploaded content whenever possible. Mention the source file/page/slide if visible in the context.
"""
    })
    return ai_chat(messages, temperature=0.2)

def render_analysis_panel(context):
    if "document_analysis" not in st.session_state:
        st.session_state.document_analysis = ""

    if context and not st.session_state.document_analysis:
        with st.spinner("SHIFA is building the document intelligence brief..."):
            st.session_state.document_analysis = generate_summary(context)

    if st.session_state.document_analysis:
        st.markdown("### AI Document Analysis")
        st.markdown(st.session_state.document_analysis)

def upload_page():

    render_title("Upload Center")

    st.markdown("""
    <div class="upload-hero">
        <h2>AI Document Intelligence Workspace</h2>
        <p>Upload notes, books, slides, text files, and images. SHIFA extracts, analyzes, remembers, and turns them into exam-ready study material.</p>
    </div>
    """, unsafe_allow_html=True)

    if "uploaded_documents" not in st.session_state:
        st.session_state.uploaded_documents = {}
    if "document_chat_history" not in st.session_state:
        st.session_state.document_chat_history = []
    if "generated_study_outputs" not in st.session_state:
        st.session_state.generated_study_outputs = {}

    uploaded_files = st.file_uploader(
        "Upload PDF, DOCX, TXT, PPTX, PNG, JPG, or JPEG files",
        type=SUPPORTED_UPLOAD_TYPES,
        accept_multiple_files=True
    )

    if st.button("🗑 Clear Uploaded Content", use_container_width=True):
        st.session_state.uploaded_documents = {}
        st.session_state.document_chat_history = []
        st.session_state.generated_study_outputs = {}
        st.session_state.document_analysis = ""
        st.rerun()

    if uploaded_files:
        progress = st.progress(0, text="Preparing upload queue...")
        new_files = []
        for uploaded_file in uploaded_files:
            file_bytes = uploaded_file.getvalue()
            digest = file_hash(file_bytes)
            if digest not in st.session_state.uploaded_documents:
                new_files.append((uploaded_file.name, file_bytes, digest))

        if new_files:
            with st.spinner("Extracting and chunking uploaded files..."):
                workers = min(4, len(new_files))
                with ThreadPoolExecutor(max_workers=workers) as executor:
                    futures = {
                        executor.submit(process_uploaded_bytes, name, file_bytes): (name, digest)
                        for name, file_bytes, digest in new_files
                    }
                    completed = 0
                    for future in as_completed(futures):
                        name, digest = futures[future]
                        completed += 1
                        try:
                            document, error = future.result()
                        except Exception as e:
                            document, error = None, f"{name} could not be processed: {e}"

                        if error:
                            st.error(error)
                        elif document:
                            st.session_state.uploaded_documents[digest] = document
                            st.session_state.document_analysis = ""

                        progress.progress(completed / len(new_files), text=f"Processed {completed} of {len(new_files)} new files")
        else:
            progress.progress(1.0, text="All selected files are already cached.")
        progress.empty()

    documents = st.session_state.uploaded_documents

    left, right = st.columns([0.36, 0.64])
    with left:
        render_file_list(documents)

    with right:
        if documents:
            render_document_metrics(documents)
            context = combined_document_context()

            preview_tabs = st.tabs(["Analysis", "Document Viewer", "Study Tools", "Chat"])

            with preview_tabs[0]:
                st.markdown('<span class="processing-ring"></span><b>Cached extraction enabled. AI analysis updates only when files change.</b>', unsafe_allow_html=True)
                render_analysis_panel(context)

            with preview_tabs[1]:
                selected_name = st.selectbox("Select document preview", [doc["name"] for doc in documents.values()])
                selected_doc = next(doc for doc in documents.values() if doc["name"] == selected_name)
                st.text_area(
                    "Extracted content preview",
                    selected_doc.get("text", "")[:12000] or "No readable text extracted.",
                    height=420
                )

            with preview_tabs[2]:
                difficulty = st.select_slider("Difficulty", options=["Easy", "Medium", "Hard"], value="Medium")
                tool_cols = st.columns(3)
                tool_map = {
                    "📄 Generate Notes": "Generate Notes",
                    "🧠 Generate Mind Map": "Generate Mind Map",
                    "📊 Generate Flowchart": "Generate Flowchart",
                    "❓ Generate MCQs": "Generate MCQs",
                    "🎯 Generate Important Questions": "Generate Important Questions",
                    "📚 Generate Revision Sheet": "Generate Revision Sheet",
                }

                for button_index, (label, task) in enumerate(tool_map.items()):
                    with tool_cols[button_index % 3]:
                        if st.button(label, use_container_width=True):
                            with st.spinner(f"{task}..."):
                                if task == "Generate MCQs":
                                    output = generate_mcqs(context, difficulty)
                                else:
                                    output = generate_notes(context, task, difficulty)
                            st.session_state.generated_study_outputs[task] = output

                for task, output in st.session_state.generated_study_outputs.items():
                    with st.expander(task, expanded=True):
                        st.markdown(output)

            with preview_tabs[3]:
                st.markdown("### 💬 Ask Questions About Uploaded Content")
                examples = [
                    "Explain page 5",
                    "Summarize chapter 2",
                    "Give important formulas",
                    "Generate 16 mark answers",
                    "Generate 2 mark questions",
                    "Explain diagrams",
                    "Give viva questions",
                    "Create MCQs",
                    "What is in this image?",
                    "Convert image notes into text",
                ]
                st.caption("Try: " + " • ".join(examples[:6]))

                for item in st.session_state.document_chat_history:
                    with st.chat_message(item["role"]):
                        st.markdown(item["content"])

                question = st.chat_input("Ask SHIFA about the uploaded content...")
                if question:
                    st.session_state.document_chat_history.append({"role": "user", "content": question})
                    with st.chat_message("user"):
                        st.markdown(question)
                    with st.chat_message("assistant"):
                        with st.spinner("Answering from uploaded content..."):
                            answer = document_chat(question, context, st.session_state.document_chat_history)
                        st.markdown(answer)
                    st.session_state.document_chat_history.append({"role": "assistant", "content": answer})
                    save_history(f"[Document Chat] {question}", answer)
                    st.rerun()
        else:
            st.info("Upload one or more files to start AI document analysis.")

# =========================================================
# MCQ PAGE
# =========================================================

def mcq_page():

    render_title("MCQ Generator")

    topic = st.text_input(
        "Enter Topic",
        placeholder="Machine Learning"
    )

    num_questions = st.slider(
        "Number of Questions",
        1,
        20,
        5
    )

    difficulty = st.select_slider(
        "Difficulty",
        options=["Easy", "Medium", "Hard"]
    )

    if st.button("Generate MCQs"):

        prompt = f"""
Generate {num_questions} {difficulty} MCQs about {topic}.

Format:
Question
A)
B)
C)
D)

Then provide answer key.
"""

        with st.spinner("Generating MCQs..."):

            answer = ai_chat([
                {
                    "role": "user",
                    "content": prompt
                }
            ])

        st.markdown(f"""
        <div class="chat-ai chat-bubble-model">
        {esc(answer)}
        </div>
        """, unsafe_allow_html=True)

# =========================================================
# HISTORY PAGE
# =========================================================

def history_page():

    render_title("Search History")

    search = st.text_input(
        "Search History",
        placeholder="Search previous chats..."
    )

    rows = get_history(search)

    if st.button("🗑 Clear History"):
        clear_history()
        st.rerun()

    if not rows:

        st.markdown("""
        <div class="history-card empty-history">
        <h4>No Search History Found</h4>
        <p>Your previous AI conversations will appear here.</p>
        </div>
        """, unsafe_allow_html=True)

        return

    for row in rows:

        st.markdown(f"""
        <div class="history-card">
        <h4>{esc(row["question"])}</h4>

        <p>{esc(row["answer"][:500])}</p>

        <div class="history-time">
        {esc(row["time"])}
        </div>
        </div>
        """, unsafe_allow_html=True)

# =========================================================
# FOOTER
# =========================================================

def footer():

    st.markdown("""
    <hr>

    <center>

    <h3 style="color:#67e8f9;">
    🚀 Built by Shifaya Simnas
    </h3>

    <p style="color:#94a3b8;">
    AI/ML Engineering Student | Python Developer | AI Builder
    </p>

    </center>
    """, unsafe_allow_html=True)

# =========================================================
# MAIN
# =========================================================

def main():

    inject_theme()

    init_db()

    if "messages" not in st.session_state:
        st.session_state.messages = []

    menu = sidebar()

    if menu == "Dashboard":
        dashboard()

    elif menu == "AI Study Chat":
        chat_page()

    elif menu == "Upload Center":
        upload_page()

    elif menu == "MCQ Generator":
        mcq_page()

    elif menu == "Search History":
        history_page()

    footer()

# =========================================================

if __name__ == "__main__":
    main()
